import os
import openpyxl
from datetime import datetime, date
from decimal import Decimal
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone

from apps.catalog.models import Color, Thickness, Product, ProductVariant, ProductStock
from apps.assets.models import MachineType, Machine, SparePart, MachineTypeSparePartCompatibility, SparePartInventoryTransaction, UtilityEquipment
from apps.inventory.models import StorageLocation, RawMaterialType, RawMaterialVariant, RawMaterialStock, RawMaterialInventoryTransaction
from apps.store.models import FinishedProductBag, StoreMovement
from apps.workforce.models import Employee, AdditionalPaymentType, PayrollConfiguration
from apps.documents.models import DocumentRegistry


class Command(BaseCommand):
    help = "Idempotently import operational factory data from ali bori NEW shoe lace.xlsx into PostgreSQL"

    def add_arguments(self, parser):
        parser.add_argument(
            'excel_file',
            nargs='?',
            default='ali bori NEW shoe lace.xlsx',
            help='Path to the Excel file'
        )

    def handle(self, *args, **options):
        excel_path = options['excel_file']
        if not os.path.exists(excel_path):
            alt_path = os.path.join('..', excel_path)
            if os.path.exists(alt_path):
                excel_path = alt_path
            else:
                self.stderr.write(self.style.ERROR(f"File not found: {excel_path}"))
                return

        self.stdout.write(self.style.SUCCESS(f"Loading workbook: {excel_path}"))
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet_map = {name.strip(): name for name in wb.sheetnames}

        summary = {
            'colors': 0,
            'thicknesses': 0,
            'products': 0,
            'product_variants': 0,
            'machines': 0,
            'machine_types': 0,
            'spare_parts': 0,
            'raw_materials': 0,
            'laborers': 0,
            'utilities': 0,
            'additional_payments': 0,
            'warnings': 0,
            'errors': 0,
        }

        color_aliases = {
            'b': ('Black', 'B', '#1A1A1A'),
            'w': ('White', 'W', '#F5F5F5'),
            '-': ('Natural/Default', '-', '#D8C39E'),
            'black': ('Black', 'B', '#1A1A1A'),
            'white': ('White', 'W', '#F5F5F5'),
            'brown': ('Brown', 'BR', '#8B461E'),
            'red': ('Red', 'RD', '#A33327'),
            'gray': ('Gray', 'GY', '#606A6D'),
        }

        def get_or_create_color(raw_name):
            if not raw_name:
                return None
            cleaned = str(raw_name).strip().lower()
            if cleaned in color_aliases:
                name, code, hex_code = color_aliases[cleaned]
            else:
                name, code, hex_code = raw_name.strip().title(), raw_name.strip()[:5].upper(), '#888888'
            color_obj, created = Color.objects.get_or_create(
                name=name,
                defaults={'code': code, 'hex_code': hex_code}
            )
            if created:
                summary['colors'] += 1
            return color_obj

        def get_or_create_thickness(raw_thick):
            if not raw_thick:
                return None
            thick_str = str(raw_thick).strip().upper()
            val_mm = None
            try:
                num_part = thick_str.replace('MM', '').strip()
                val_mm = Decimal(num_part)
            except Exception:
                pass
            t_obj, created = Thickness.objects.get_or_create(
                name=thick_str,
                defaults={'value_mm': val_mm}
            )
            if created:
                summary['thicknesses'] += 1
            return t_obj

        # 1. PRODUCTES SHEET
        if 'PRODUCTES' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing PRODUCTES...")
                ws = wb[sheet_map['PRODUCTES']]
                rows = list(ws.iter_rows(values_only=True))
                shoe_lace_product, p_created = Product.objects.get_or_create(
                    code="SHOE-LACE",
                    defaults={'name': 'Standard Braided & Tipped Shoe Lace', 'category': 'SHOE_LACE', 'unit': 'KG'}
                )
                if p_created:
                    summary['products'] += 1

                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    serial_code = str(r[0]).strip()
                    color_val = r[1]
                    thick_val = r[2]
                    daily_prod = Decimal(str(r[3] or 0)) if r[3] is not None else Decimal('0.00')
                    in_qty = Decimal(str(r[4] or 0)) if r[4] is not None else Decimal('0.00')
                    out_qty = Decimal(str(r[5] or 0)) if r[5] is not None else Decimal('0.00')
                    st_v = Decimal(str(r[6] or 0)) if r[6] is not None else Decimal('0.00')
                    total = Decimal(str(r[7] or 0)) if r[7] is not None else Decimal('0.00')
                    remark = str(r[8]).strip() if r[8] is not None else ""

                    color_obj = get_or_create_color(color_val)
                    thick_obj = get_or_create_thickness(thick_val)

                    variant, v_created = ProductVariant.objects.get_or_create(
                        serial_code=serial_code,
                        color=color_obj,
                        thickness=thick_obj,
                        defaults={'product': shoe_lace_product}
                    )
                    if v_created:
                        summary['product_variants'] += 1

                    ProductStock.objects.update_or_create(
                        variant=variant,
                        defaults={
                            'daily_product': daily_prod,
                            'in_qty': in_qty,
                            'out_qty': out_qty,
                            'st_v': st_v,
                            'total': total,
                            'calculated_stock': total,
                            'remark': remark,
                        }
                    )

        # 2. MACHINES SHEET (Batch optimized)
        if 'MACHINES' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing MACHINES...")
                ws = wb[sheet_map['MACHINES']]
                rows = list(ws.iter_rows(values_only=True))
                known_types = ['TIPPING', 'SPINDLE', 'CH SPINDLE', 'WINDER', 'SMALL WINDER']
                type_objs = {}
                for kt in known_types:
                    m_type, t_created = MachineType.objects.get_or_create(
                        name=kt,
                        defaults={'description': f'Machine type: {kt}', 'maintenance_interval_days': 30}
                    )
                    type_objs[kt] = m_type
                    if t_created:
                        summary['machine_types'] += 1

                # Deduplicate sheet rows by machine_code (favoring rows with more non-empty data)
                machine_data_by_code = {}
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    m_code = str(r[0]).strip()
                    if not m_code:
                        continue
                    if m_code in machine_data_by_code:
                        prev_r = machine_data_by_code[m_code]
                        prev_non_empty = sum(1 for x in prev_r if x is not None and str(x).strip() != '')
                        curr_non_empty = sum(1 for x in r if x is not None and str(x).strip() != '')
                        if curr_non_empty > prev_non_empty:
                            machine_data_by_code[m_code] = r
                    else:
                        machine_data_by_code[m_code] = r

                existing_machines = {m.machine_code: m for m in Machine.objects.all()}
                machines_to_create = []
                machines_to_update = []

                for m_code, r in machine_data_by_code.items():
                    upper_code = m_code.upper()
                    if upper_code.startswith('TP'):
                        chosen_type = type_objs.get('TIPPING')
                    elif upper_code.startswith('SP'):
                        chosen_type = type_objs.get('SPINDLE')
                    elif upper_code.startswith('CH'):
                        chosen_type = type_objs.get('CH SPINDLE')
                    elif upper_code.startswith('SW') or 'SMALL' in upper_code:
                        chosen_type = type_objs.get('SMALL WINDER')
                    elif upper_code.startswith('WD') or 'WINDER' in upper_code:
                        chosen_type = type_objs.get('WINDER')
                    else:
                        chosen_type = type_objs.get('TIPPING')

                    room = str(r[2]).strip() if r[2] is not None else "B"
                    prod_type = str(r[3]).strip() if r[3] is not None else "CUT"
                    place = str(r[4]).strip() if r[4] is not None else "B1"
                    health_raw = str(r[5]).strip().upper() if r[5] is not None else "NORMAL"
                    health = Machine.Health.NORMAL
                    if "SERVICE" in health_raw:
                        health = Machine.Health.NEEDS_SERVICE
                    elif "CRITICAL" in health_raw or "BAD" in health_raw:
                        health = Machine.Health.CRITICAL

                    most_issue = str(r[6]).strip() if r[6] is not None else ""
                    service_time = r[7] if isinstance(r[7], (datetime, date)) else None
                    if isinstance(service_time, datetime) and timezone.is_naive(service_time):
                        service_time = timezone.make_aware(service_time)

                    perf = Decimal('1.00')
                    try:
                        if r[8] is not None:
                            val = Decimal(str(r[8]))
                            if val > Decimal('1.00') and val <= Decimal('100.00'):
                                val = val / Decimal('100.00')
                            perf = max(Decimal('0.00'), min(Decimal('1.00'), val))
                    except Exception:
                        pass

                    total_count = 1
                    try:
                        if r[9] is not None:
                            total_count = int(r[9])
                    except Exception:
                        pass

                    remark = str(r[10]).strip() if len(r) > 10 and r[10] is not None else ""

                    if m_code in existing_machines:
                        m_obj = existing_machines[m_code]
                        m_obj.machine_type = chosen_type
                        m_obj.building = "B1" if "1" in place else "B2"
                        m_obj.room = room
                        m_obj.product_type = prod_type
                        m_obj.place = place
                        m_obj.health = health
                        m_obj.most_frequent_issue = most_issue
                        m_obj.service_time = service_time if isinstance(service_time, datetime) else None
                        m_obj.last_service_date = service_time.date() if isinstance(service_time, datetime) else (service_time if isinstance(service_time, date) else None)
                        m_obj.performance_score = perf
                        m_obj.total = total_count
                        m_obj.remark = remark
                        machines_to_update.append(m_obj)
                    else:
                        m_obj = Machine(
                            machine_code=m_code,
                            name=f"{chosen_type.name} - {m_code}",
                            machine_type=chosen_type,
                            building="B1" if "1" in place else "B2",
                            room=room,
                            product_type=prod_type,
                            place=place,
                            health=health,
                            most_frequent_issue=most_issue,
                            service_time=service_time if isinstance(service_time, datetime) else None,
                            last_service_date=service_time.date() if isinstance(service_time, datetime) else (service_time if isinstance(service_time, date) else None),
                            performance_score=perf,
                            total=total_count,
                            remark=remark,
                        )
                        machines_to_create.append(m_obj)

                if machines_to_create:
                    Machine.objects.bulk_create(machines_to_create, ignore_conflicts=True)
                    summary['machines'] += len(machines_to_create)
                if machines_to_update:
                    Machine.objects.bulk_update(
                        machines_to_update,
                        ['machine_type', 'building', 'room', 'product_type', 'place', 'health', 'most_frequent_issue', 'service_time', 'last_service_date', 'performance_score', 'total', 'remark']
                    )
                    summary['machines'] += len(machines_to_update)

        # 3. MACHINE PARTS SHEET
        if 'MACHINE PARTS' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing MACHINE PARTS compatibility...")
                ws = wb[sheet_map['MACHINE PARTS']]
                rows = list(ws.iter_rows(values_only=True))
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    m_type_name = str(r[0]).strip().upper()
                    m_type_obj = MachineType.objects.filter(name__iexact=m_type_name).first()
                    if not m_type_obj:
                        m_type_obj = MachineType.objects.create(name=m_type_name)

                    part_cols = [r[1], r[2], r[3], r[4], r[5]]
                    for idx, part_name in enumerate(part_cols):
                        if not part_name or str(part_name).strip().upper() in ['N', 'NONE', '-', '']:
                            continue
                        clean_part = str(part_name).strip()
                        part_code = clean_part.upper().replace(' ', '_')
                        spare_obj, sp_created = SparePart.objects.get_or_create(
                            part_code=part_code,
                            defaults={'name': clean_part, 'quantity': 0, 'minimum_stock': 3}
                        )
                        if sp_created:
                            summary['spare_parts'] += 1
                        spare_obj.compatible_machine_types.add(m_type_obj)
                        MachineTypeSparePartCompatibility.objects.get_or_create(
                            machine_type=m_type_obj,
                            spare_part=spare_obj,
                            defaults={'part_slot': f"PART NO{idx+1}"}
                        )

        # 4. SPARES SHEET
        if 'SPARES' in sheet_map or 'SPARES ' in sheet_map:
            with transaction.atomic():
                s_key = 'SPARES' if 'SPARES' in sheet_map else 'SPARES '
                self.stdout.write(f"Importing {s_key}...")
                ws = wb[sheet_map[s_key]]
                rows = list(ws.iter_rows(values_only=True))
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    p_name = str(r[0]).strip()
                    part_code = p_name.upper().replace(' ', '_')
                    for_m = str(r[1]).strip() if r[1] is not None else ""
                    place = str(r[2]).strip() if r[2] is not None else ""
                    room = str(r[3]).strip() if r[3] is not None else ""
                    shelf = str(r[4]).strip() if r[4] is not None else ""
                    qty = Decimal(str(r[5] or 0)) if r[5] is not None else Decimal('0.00')
                    rem = str(r[6]).strip() if len(r) > 6 and r[6] is not None else ""

                    sp_obj, sp_created = SparePart.objects.update_or_create(
                        part_code=part_code,
                        defaults={
                            'name': p_name,
                            'for_machine_text': for_m,
                            'place': place,
                            'room': room,
                            'shelf': shelf,
                            'quantity': qty,
                            'minimum_stock': 3.00,
                            'remark': rem,
                        }
                    )
                    if sp_created:
                        summary['spare_parts'] += 1

                    if qty > 0 and not SparePartInventoryTransaction.objects.filter(spare_part=sp_obj, transaction_type='INITIAL_STOCK').exists():
                        SparePartInventoryTransaction.objects.create(
                            spare_part=sp_obj,
                            transaction_type=SparePartInventoryTransaction.TransactionType.INITIAL_STOCK,
                            quantity=qty,
                            balance_after=qty,
                            notes="Imported opening stock from Excel SPARES sheet"
                        )

        # 5. RAW MATERIALS SHEET
        if 'RAW MATERIALS' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing RAW MATERIALS...")
                ws = wb[sheet_map['RAW MATERIALS']]
                rows = list(ws.iter_rows(values_only=True))
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    m_name = str(r[0]).strip()
                    color_raw = str(r[1]).strip() if r[1] is not None else "General"
                    st_v = Decimal(str(r[2] or 0)) if r[2] is not None else Decimal('0.00')
                    st_n = Decimal(str(r[3] or 0)) if r[3] is not None else Decimal('0.00')
                    place_str = str(r[4]).strip() if r[4] is not None else "ST 1"
                    total_kg = Decimal(str(r[5] or 0)) if r[5] is not None else Decimal('0.00')

                    mat_type, _ = RawMaterialType.objects.get_or_create(
                        name=m_name,
                        defaults={'code': m_name.upper().replace(' ', '_')[:30]}
                    )
                    color_obj = get_or_create_color(color_raw)
                    v_code = f"{mat_type.code}-{color_raw.upper()[:4]}"

                    rm_variant, v_created = RawMaterialVariant.objects.get_or_create(
                        material_type=mat_type,
                        color_name=color_raw.lower(),
                        defaults={'code': v_code, 'color': color_obj, 'minimum_stock_kg': Decimal('25.00')}
                    )
                    if v_created:
                        summary['raw_materials'] += 1

                    loc_obj, _ = StorageLocation.objects.get_or_create(
                        code=place_str,
                        defaults={'name': f"Storage Location {place_str}"}
                    )

                    RawMaterialStock.objects.update_or_create(
                        variant=rm_variant,
                        location=loc_obj,
                        defaults={
                            'place_text': place_str,
                            'st_v': st_v,
                            'st_n': st_n,
                            'total_kg': total_kg,
                            'available_kg': total_kg,
                        }
                    )

                    if total_kg > 0 and not RawMaterialInventoryTransaction.objects.filter(variant=rm_variant, transaction_type='INITIAL_STOCK').exists():
                        RawMaterialInventoryTransaction.objects.create(
                            variant=rm_variant,
                            transaction_type=RawMaterialInventoryTransaction.TransactionType.INITIAL_STOCK,
                            quantity_kg=total_kg,
                            balance_after_kg=total_kg,
                            notes=f"Imported opening stock ({place_str})"
                        )

        # 6. LABORERS SHEET
        if 'LABORERS' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing LABORERS...")
                ws = wb[sheet_map['LABORERS']]
                rows = list(ws.iter_rows(values_only=True))
                for idx, r in enumerate(rows[1:], start=1):
                    if not r or not r[0]:
                        continue
                    l_name = str(r[0]).strip()
                    age = int(r[1]) if r[1] is not None else None
                    gender = str(r[2]).strip().upper() if r[2] is not None else "M"
                    if gender not in ['M', 'F']:
                        gender = "M"
                    work_hour = int(r[3]) if r[3] is not None else 8
                    work_room = str(r[4]).strip() if r[4] is not None else "B1"
                    salary = Decimal(str(r[5] or 0)) if r[5] is not None else Decimal('0.00')
                    perf = Decimal(str(r[6] or 80)) if r[6] is not None else Decimal('80.00')
                    remark = str(r[7]).strip() if len(r) > 7 and r[7] is not None else ""

                    emp_id = f"EMP-{idx:04d}"
                    emp_obj, e_created = Employee.objects.update_or_create(
                        name=l_name,
                        defaults={
                            'employee_id': emp_id,
                            'age': age,
                            'gender': gender,
                            'work_hours_per_day': work_hour,
                            'work_room': work_room,
                            'base_salary': salary,
                            'performance_score': perf,
                            'remark': remark,
                        }
                    )
                    if e_created:
                        summary['laborers'] += 1

        # 7. UTILITIES SHEET
        if 'UTILITIES' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing UTILITIES...")
                ws = wb[sheet_map['UTILITIES']]
                rows = list(ws.iter_rows(values_only=True))
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    u_name = str(r[0]).strip()
                    u_room = str(r[1]).strip() if r[1] is not None else "B"
                    qty = int(r[2]) if r[2] is not None else 1

                    u_obj, u_created = UtilityEquipment.objects.update_or_create(
                        name=u_name,
                        room=u_room,
                        defaults={'quantity': qty}
                    )
                    if u_created:
                        summary['utilities'] += 1

        # 8. ADDITIONAL PAYMENTS SHEET
        if 'ADDITIONAL PAYMENTS' in sheet_map:
            with transaction.atomic():
                self.stdout.write("Importing ADDITIONAL PAYMENTS...")
                ws = wb[sheet_map['ADDITIONAL PAYMENTS']]
                rows = list(ws.iter_rows(values_only=True))
                for r in rows[1:]:
                    if not r or not r[0]:
                        continue
                    pay_name = str(r[0]).strip()
                    means = str(r[1]).strip() if r[1] is not None else "WEEKLY"
                    for_who = str(r[2]).strip() if r[2] is not None else "F.I.A"
                    price = Decimal(str(r[3] or 1800)) if r[3] is not None else Decimal('1800.00')

                    pay_type, p_created = AdditionalPaymentType.objects.update_or_create(
                        name=pay_name,
                        defaults={
                            'means': means,
                            'target_group': for_who,
                            'default_price': price,
                        }
                    )
                    if p_created:
                        summary['additional_payments'] += 1

        # 9. Default Payroll Configuration
        if not PayrollConfiguration.objects.exists():
            PayrollConfiguration.objects.create(
                name="Factory Standard Monthly Policy",
                salary_period=PayrollConfiguration.SalaryPeriod.MONTHLY,
                working_days_per_period=26,
                absence_deduction_method=PayrollConfiguration.AbsenceDeductionMethod.DAILY_RATE,
                absence_deduction_rate=Decimal('0.0385'),
                overtime_hourly_multiplier=Decimal('1.25'),
                saturday_rate=Decimal('1800.00'),
                is_active=True
            )

        # Output Summary (ASCII characters only for safe Windows cp1252 output)
        self.stdout.write(self.style.SUCCESS("\n=========================================="))
        self.stdout.write(self.style.SUCCESS("[OK] FACTORY DATA IMPORT COMPLETED SUCCESSFULLY"))
        self.stdout.write(self.style.SUCCESS("=========================================="))
        self.stdout.write(f"  Products:            {summary['products']}")
        self.stdout.write(f"  Product Variants:    {summary['product_variants']}")
        self.stdout.write(f"  Colors:              {summary['colors']}")
        self.stdout.write(f"  Thicknesses:         {summary['thicknesses']}")
        self.stdout.write(f"  Machines:            {summary['machines']}")
        self.stdout.write(f"  Machine Types:       {summary['machine_types']}")
        self.stdout.write(f"  Spare Parts:         {summary['spare_parts']}")
        self.stdout.write(f"  Raw Materials:       {summary['raw_materials']}")
        self.stdout.write(f"  Laborers/Employees:  {summary['laborers']}")
        self.stdout.write(f"  Utilities:           {summary['utilities']}")
        self.stdout.write(f"  Additional Payments: {summary['additional_payments']}")
        self.stdout.write(self.style.SUCCESS("==========================================\n"))
